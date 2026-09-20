import os
import io
import re
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference

EXCEL_COLUMNS = [
    "Job ID",
    "Priority",
    "Company",
    "Job Role",
    "Function / Category",
    "Location",
    "Work Mode",
    "Experience / Eligibility",
    "Application Date",
    "Application Status",
    "Interview",
    "Offer",
    "Assessment"
]

CATEGORIES = [
    "Data Analytics", "Business Analytics", "Finance", "Accounting", 
    "HR", "Business Intelligence", "IT", "Consulting", "MIS", "Other"
]

PRIORITIES = ["High", "Medium", "Low"]
WORK_MODES = ["On-site", "Hybrid", "Remote", "Not Mentioned"]
APPLICATION_STATUSES = [
    "Saved", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn", "On Hold"
]
YES_NO = ["Yes", "No"]

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data"))
WORKBOOK_FILENAME = "AI_Job_Tracker.xlsx"
WORKBOOK_PATH = os.path.join(DATA_DIR, WORKBOOK_FILENAME)


def get_workbook_path() -> str:
    os.makedirs(DATA_DIR, exist_ok=True)
    return WORKBOOK_PATH


def apply_header_styles(ws):
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[1].height = 28
    for col_num, col_name in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_num, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border


def apply_data_validations(ws, max_row=1000):
    # Data validations for controlled dropdown fields
    dv_priority = DataValidation(type="list", formula1=f'"{",".join(PRIORITIES)}"', allow_blank=True)
    dv_category = DataValidation(type="list", formula1=f'"{",".join(CATEGORIES)}"', allow_blank=True)
    dv_work_mode = DataValidation(type="list", formula1=f'"{",".join(WORK_MODES)}"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1=f'"{",".join(APPLICATION_STATUSES)}"', allow_blank=True)
    dv_yesno1 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_yesno2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_yesno3 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)

    ws.add_data_validation(dv_priority)
    ws.add_data_validation(dv_category)
    ws.add_data_validation(dv_work_mode)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_yesno1)
    ws.add_data_validation(dv_yesno2)
    ws.add_data_validation(dv_yesno3)

    dv_priority.add(f"B2:B{max_row}")
    dv_category.add(f"E2:E{max_row}")
    dv_work_mode.add(f"G2:G{max_row}")
    dv_status.add(f"J2:J{max_row}")
    dv_yesno1.add(f"K2:K{max_row}")
    dv_yesno2.add(f"L2:L{max_row}")
    dv_yesno3.add(f"M2:M{max_row}")


def format_job_tracker_sheet(ws):
    apply_header_styles(ws)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 2)}"
    apply_data_validations(ws)

    # Column widths
    col_widths = {
        "A": 12, "B": 14, "C": 24, "D": 26, "E": 22,
        "F": 18, "G": 16, "H": 22, "I": 18, "J": 18,
        "K": 12, "L": 12, "M": 14
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width


def update_summary_sheet(wb):
    if "Summary" in wb.sheetnames:
        del wb["Summary"]

    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    section_font = Font(name="Calibri", size=12, bold=True, color="334155")
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    kpi_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    kpi_font = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
    kpi_label_font = Font(name="Calibri", size=10, bold=True, color="64748B")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title
    ws_sum["A1"] = "AI Job Tracker - Summary & Analytics"
    ws_sum["A1"].font = title_font

    # KPI Section
    ws_sum["A3"] = "TOTAL JOBS"
    ws_sum["A3"].font = kpi_label_font
    ws_sum["A4"] = "='Job Tracker'!A2:A1000"  # formula placeholder or counta
    ws_sum["A4"].value = "=COUNTA('Job Tracker'!A2:A1000)"
    ws_sum["A4"].font = Font(name="Calibri", size=22, bold=True, color="1E3A8A")

    ws_sum["C3"] = "APPLIED"
    ws_sum["C3"].font = kpi_label_font
    ws_sum["C4"] = '=COUNTIF(\'Job Tracker\'!J2:J1000, "Applied")'
    ws_sum["C4"].font = Font(name="Calibri", size=22, bold=True, color="0284C7")

    ws_sum["E3"] = "ASSESSMENTS"
    ws_sum["E3"].font = kpi_label_font
    ws_sum["E4"] = '=COUNTIF(\'Job Tracker\'!M2:M1000, "Yes")'
    ws_sum["E4"].font = Font(name="Calibri", size=22, bold=True, color="D97706")

    ws_sum["G3"] = "INTERVIEWS"
    ws_sum["G3"].font = kpi_label_font
    ws_sum["G4"] = '=COUNTIF(\'Job Tracker\'!K2:K1000, "Yes")'
    ws_sum["G4"].font = Font(name="Calibri", size=22, bold=True, color="7C3AED")

    ws_sum["I3"] = "OFFERS"
    ws_sum["I3"].font = kpi_label_font
    ws_sum["I4"] = '=COUNTIF(\'Job Tracker\'!L2:L1000, "Yes")'
    ws_sum["I4"].font = Font(name="Calibri", size=22, bold=True, color="16A34A")

    # Application Status Table
    ws_sum["A7"] = "Application Status Summary"
    ws_sum["A7"].font = section_font
    ws_sum["A8"] = "Status"
    ws_sum["B8"] = "Count"
    ws_sum["A8"].fill = header_fill
    ws_sum["B8"].fill = header_fill
    ws_sum["A8"].font = Font(bold=True)
    ws_sum["B8"].font = Font(bold=True)

    row = 9
    for status in APPLICATION_STATUSES:
        ws_sum.cell(row=row, column=1, value=status).border = thin_border
        cell_val = ws_sum.cell(row=row, column=2, value=f'=COUNTIF(\'Job Tracker\'!J2:J1000, "{status}")')
        cell_val.border = thin_border
        row += 1

    # Location Table (Dynamic city list from Job Tracker)
    ws_sum["D7"] = "Location Breakdown (City Level)"
    ws_sum["D7"].font = section_font
    ws_sum["D8"] = "Location"
    ws_sum["E8"] = "Count"
    ws_sum["D8"].fill = header_fill
    ws_sum["E8"].fill = header_fill
    ws_sum["D8"].font = Font(bold=True)
    ws_sum["E8"].font = Font(bold=True)

    # Extract distinct locations from Job Tracker sheet
    ws_jt = wb["Job Tracker"]
    locations = set()
    for r in range(2, ws_jt.max_row + 1):
        loc = ws_jt.cell(row=r, column=6).value
        if loc and str(loc).strip():
            locations.add(str(loc).strip())
    
    sorted_locations = sorted(list(locations)) if locations else ["Bangalore", "Chennai", "Hyderabad", "Mumbai", "Pune", "Delhi", "Kochi", "Remote"]

    loc_start_row = 9
    for i, loc in enumerate(sorted_locations):
        curr_row = loc_start_row + i
        ws_sum.cell(row=curr_row, column=4, value=loc).border = thin_border
        cell_loc = ws_sum.cell(row=curr_row, column=5, value=f'=COUNTIF(\'Job Tracker\'!F2:F1000, "{loc}")')
        cell_loc.border = thin_border

    loc_end_row = loc_start_row + len(sorted_locations) - 1

    # Add Bar Chart for Locations
    if len(sorted_locations) > 0:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Jobs by Location"
        chart.y_axis.title = "City Location"
        chart.x_axis.title = "Job Count"

        data = Reference(ws_sum, min_col=5, min_row=8, max_row=loc_end_row)
        cats = Reference(ws_sum, min_col=4, min_row=9, max_row=loc_end_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        chart.width = 16
        chart.height = 10
        ws_sum.add_chart(chart, "G7")

    # Column widths
    ws_sum.column_dimensions['A'].width = 22
    ws_sum.column_dimensions['B'].width = 12
    ws_sum.column_dimensions['C'].width = 14
    ws_sum.column_dimensions['D'].width = 22
    ws_sum.column_dimensions['E'].width = 12
    ws_sum.column_dimensions['F'].width = 4


def init_tracker_workbook() -> openpyxl.Workbook:
    path = get_workbook_path()
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        return wb

    wb = openpyxl.Workbook()
    # Sheet 1: Job Tracker
    ws_jt = wb.active
    ws_jt.title = "Job Tracker"
    format_job_tracker_sheet(ws_jt)

    # Sheet 2: Summary
    update_summary_sheet(wb)

    wb.save(path)
    return wb


def get_all_jobs() -> List[Dict[str, Any]]:
    path = get_workbook_path()
    if not os.path.exists(path):
        init_tracker_workbook()
    
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Job Tracker" not in wb.sheetnames:
        return []
    
    ws = wb["Job Tracker"]
    jobs = []
    
    # Header check
    headers = [ws.cell(row=1, column=col).value for col in range(1, 14)]
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for r in range(2, ws.max_row + 1):
        job_id = ws.cell(row=r, column=1).value
        if not job_id or not str(job_id).strip():
            continue
        
        job = {
            "job_id": str(ws.cell(row=r, column=1).value or "").strip(),
            "priority": str(ws.cell(row=r, column=2).value or "Medium").strip(),
            "company": str(ws.cell(row=r, column=3).value or "").strip(),
            "job_role": str(ws.cell(row=r, column=4).value or "").strip(),
            "function_category": str(ws.cell(row=r, column=5).value or "Other").strip(),
            "location": str(ws.cell(row=r, column=6).value or "").strip(),
            "work_mode": str(ws.cell(row=r, column=7).value or "Not Mentioned").strip(),
            "experience_eligibility": str(ws.cell(row=r, column=8).value or "Not Mentioned").strip(),
            "application_date": str(ws.cell(row=r, column=9).value or "").strip(),
            "application_status": str(ws.cell(row=r, column=10).value or "Saved").strip(),
            "interview": str(ws.cell(row=r, column=11).value or "No").strip(),
            "offer": str(ws.cell(row=r, column=12).value or "No").strip(),
            "assessment": str(ws.cell(row=r, column=13).value or "No").strip(),
        }
        jobs.append(job)
        
    wb.close()
    return jobs


def get_next_job_id() -> str:
    jobs = get_all_jobs()
    if not jobs:
        return "J001"
    
    max_num = 0
    for job in jobs:
        jid = job.get("job_id", "")
        match = re.search(r"J(\d+)", jid, re.IGNORECASE)
        if match:
            num = int(match.group(1))
            if num > max_num:
                max_num = num
    
    next_num = max_num + 1
    return f"J{next_num:03d}"


def add_job(job_data: Dict[str, Any]) -> Dict[str, Any]:
    path = get_workbook_path()
    if not os.path.exists(path):
        init_tracker_workbook()
        
    wb = openpyxl.load_workbook(path)
    ws_jt = wb["Job Tracker"]
    
    if not job_data.get("job_id"):
        job_data["job_id"] = get_next_job_id()

    new_row = ws_jt.max_row + 1
    # Handle case if row 2 is empty header row
    if new_row == 2 and ws_jt.cell(row=2, column=1).value is None:
        new_row = 2

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    values = [
        job_data.get("job_id"),
        job_data.get("priority", "Medium"),
        job_data.get("company", ""),
        job_data.get("job_role", ""),
        job_data.get("function_category", "Other"),
        job_data.get("location", ""),
        job_data.get("work_mode", "Not Mentioned"),
        job_data.get("experience_eligibility", "Not Mentioned"),
        job_data.get("application_date", ""),
        job_data.get("application_status", "Saved"),
        job_data.get("interview", "No"),
        job_data.get("offer", "No"),
        job_data.get("assessment", "No")
    ]

    for col_idx, val in enumerate(values, 1):
        cell = ws_jt.cell(row=new_row, column=col_idx, value=val)
        cell.border = thin_border
        if col_idx in [1, 2, 7, 9, 10, 11, 12, 13]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

    format_job_tracker_sheet(ws_jt)
    update_summary_sheet(wb)
    
    wb.save(path)
    wb.close()
    return job_data


def update_job(job_id: str, updates: Dict[str, Any]) -> Tuple[bool, Optional[Dict[str, Any]]]:
    path = get_workbook_path()
    if not os.path.exists(path):
        return False, None
        
    wb = openpyxl.load_workbook(path)
    ws_jt = wb["Job Tracker"]
    
    target_row = None
    for r in range(2, ws_jt.max_row + 1):
        val = ws_jt.cell(row=r, column=1).value
        if val and str(val).strip().upper() == str(job_id).strip().upper():
            target_row = r
            break
            
    if not target_row:
        wb.close()
        return False, None

    col_map = {
        "job_id": 1,
        "priority": 2,
        "company": 3,
        "job_role": 4,
        "function_category": 5,
        "location": 6,
        "work_mode": 7,
        "experience_eligibility": 8,
        "application_date": 9,
        "application_status": 10,
        "interview": 11,
        "offer": 12,
        "assessment": 13
    }

    for key, val in updates.items():
        if key in col_map:
            col_idx = col_map[key]
            ws_jt.cell(row=target_row, column=col_idx, value=val)

    update_summary_sheet(wb)
    wb.save(path)
    
    # Read updated job data
    updated_job = {}
    for key, col_idx in col_map.items():
        updated_job[key] = str(ws_jt.cell(row=target_row, column=col_idx).value or "").strip()
        
    wb.close()
    return True, updated_job


def delete_job(job_id: str) -> bool:
    path = get_workbook_path()
    if not os.path.exists(path):
        return False
        
    wb = openpyxl.load_workbook(path)
    ws_jt = wb["Job Tracker"]
    
    target_row = None
    for r in range(2, ws_jt.max_row + 1):
        val = ws_jt.cell(row=r, column=1).value
        if val and str(val).strip().upper() == str(job_id).strip().upper():
            target_row = r
            break
            
    if not target_row:
        wb.close()
        return False

    ws_jt.delete_rows(target_row)
    update_summary_sheet(wb)
    wb.save(path)
    wb.close()
    return True


def preview_import_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    try:
        wb_import = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        ws = wb_import.active
        
        # Read header row
        headers = []
        for col in range(1, ws.max_column + 1):
            h_val = ws.cell(row=1, column=col).value
            if h_val:
                headers.append(str(h_val).strip())

        # Match columns with standard 13
        matched_columns = []
        ignored_columns = []

        standard_lookup = {col.lower().replace(" ", "").replace("/", ""): col for col in EXCEL_COLUMNS}

        for h in headers:
            normalized = h.lower().replace(" ", "").replace("/", "").replace("_", "")
            found = None
            for key, val in standard_lookup.items():
                if key in normalized or normalized in key:
                    found = val
                    break
            if found:
                matched_columns.append({"source": h, "target": found})
            else:
                ignored_columns.append(h)

        data_rows_count = 0
        for r in range(2, ws.max_row + 1):
            row_has_val = any(ws.cell(row=r, column=c).value is not None for c in range(1, len(headers) + 1))
            if row_has_val:
                data_rows_count += 1

        wb_import.close()
        return {
            "filename": filename,
            "total_jobs_found": data_rows_count,
            "columns_recognized_count": len(matched_columns),
            "columns_ignored_count": len(ignored_columns),
            "matched_columns": matched_columns,
            "ignored_columns": ignored_columns
        }
    except Exception as e:
        raise ValueError(f"Could not read Excel file: {str(e)}")


def confirm_import_file(file_bytes: bytes, mode: str) -> Dict[str, Any]:
    # Mode can be 'new' (reset AI_Job_Tracker.xlsx) or 'merge' (append non-duplicates)
    preview_info = preview_import_file(file_bytes, "uploaded.xlsx")
    wb_import = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    ws_import = wb_import.active

    # Column index map for import file
    header_indices = {}
    for col in range(1, ws_import.max_column + 1):
        h_val = ws_import.cell(row=1, column=col).value
        if h_val:
            header_indices[str(h_val).strip()] = col

    target_path = get_workbook_path()

    if mode == "new":
        if os.path.exists(target_path):
            os.remove(target_path)
        init_tracker_workbook()

    existing_jobs = get_all_jobs()
    existing_job_ids = {j["job_id"].upper() for j in existing_jobs}
    
    # Helper to map standard field to imported column value
    matched_mapping = {m["target"]: m["source"] for m in preview_info["matched_columns"]}

    imported_count = 0
    duplicate_count = 0

    for r in range(2, ws_import.max_row + 1):
        row_has_val = any(ws_import.cell(row=r, column=c).value is not None for c in range(1, ws_import.max_column + 1))
        if not row_has_val:
            continue

        def get_val(target_col_name: str, default: str = "") -> str:
            src_col = matched_mapping.get(target_col_name)
            if src_col and src_col in header_indices:
                idx = header_indices[src_col]
                v = ws_import.cell(row=r, column=idx).value
                return str(v).strip() if v is not None else default
            return default

        incoming_company = get_val("Company")
        incoming_role = get_val("Job Role")
        incoming_loc = get_val("Location")

        if not incoming_company and not incoming_role:
            continue

        # Check for duplicate if mode is merge
        if mode == "merge":
            is_dup = False
            for ej in existing_jobs:
                if (ej.get("company", "").lower() == incoming_company.lower() and
                    ej.get("job_role", "").lower() == incoming_role.lower() and
                    ej.get("location", "").lower() == incoming_loc.lower()):
                    is_dup = True
                    break
            if is_dup:
                duplicate_count += 1
                continue

        job_id = get_val("Job ID")
        if not job_id or job_id.upper() in existing_job_ids or mode == "merge":
            job_id = get_next_job_id()
            existing_job_ids.add(job_id.upper())

        job_dict = {
            "job_id": job_id,
            "priority": get_val("Priority", "Medium") or "Medium",
            "company": incoming_company or "Unknown Company",
            "job_role": incoming_role or "Role Not Specified",
            "function_category": get_val("Function / Category", "Other") or "Other",
            "location": incoming_loc or "Not Mentioned",
            "work_mode": get_val("Work Mode", "Not Mentioned") or "Not Mentioned",
            "experience_eligibility": get_val("Experience / Eligibility", "Not Mentioned") or "Not Mentioned",
            "application_date": get_val("Application Date", ""),
            "application_status": get_val("Application Status", "Saved") or "Saved",
            "interview": get_val("Interview", "No") or "No",
            "offer": get_val("Offer", "No") or "No",
            "assessment": get_val("Assessment", "No") or "No"
        }

        add_job(job_dict)
        imported_count += 1
        existing_jobs.append(job_dict)

    wb_import.close()
    return {
        "success": True,
        "mode": mode,
        "imported_count": imported_count,
        "duplicate_skipped_count": duplicate_count,
        "total_jobs_in_tracker": len(get_all_jobs())
    }


def get_summary_data() -> Dict[str, Any]:
    jobs = get_all_jobs()
    total_jobs = len(jobs)

    status_counts = {status: 0 for status in APPLICATION_STATUSES}
    priority_counts = {p: 0 for p in PRIORITIES}
    work_mode_counts = {wm: 0 for wm in WORK_MODES}
    category_counts = {cat: 0 for cat in CATEGORIES}
    location_counts: Dict[str, int] = {}

    applied_count = 0
    assessment_count = 0
    interview_count = 0
    offer_count = 0

    for job in jobs:
        status = job.get("application_status", "Saved")
        if status in status_counts:
            status_counts[status] += 1
        else:
            status_counts["Saved"] += 1

        priority = job.get("priority", "Medium")
        if priority in priority_counts:
            priority_counts[priority] += 1

        wm = job.get("work_mode", "Not Mentioned")
        if wm in work_mode_counts:
            work_mode_counts[wm] += 1

        cat = job.get("function_category", "Other")
        if cat in category_counts:
            category_counts[cat] += 1

        loc = job.get("location", "Not Mentioned")
        if loc:
            loc_key = loc.strip()
            location_counts[loc_key] = location_counts.get(loc_key, 0) + 1

        if status == "Applied":
            applied_count += 1
        if job.get("assessment", "No").lower() == "yes":
            assessment_count += 1
        if job.get("interview", "No").lower() == "yes":
            interview_count += 1
        if job.get("offer", "No").lower() == "yes":
            offer_count += 1

    # Sort locations by count descending
    sorted_locations = dict(sorted(location_counts.items(), key=lambda item: item[1], reverse=True))

    return {
        "total_jobs": total_jobs,
        "applied": applied_count,
        "assessments": assessment_count,
        "interviews": interview_count,
        "offers": offer_count,
        "status_breakdown": status_counts,
        "priority_breakdown": priority_counts,
        "work_mode_breakdown": work_mode_counts,
        "category_breakdown": category_counts,
        "location_breakdown": sorted_locations
    }
